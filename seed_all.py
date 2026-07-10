"""
seed_all.py
Comprehensive database seeder for lps_system.
Run: .\.venv\Scripts\python.exe ..\seed_all.py
(from backend/ directory, or adjust sys.path below)

Seeds:
  1. All-role users
  2. Production lines
  3. Vehicle models (CarModel) from Excel
  4. Master data (parts/BOM) from Excel
"""

import os
import sys
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'backend'))

from app import create_app
from app.extensions import db
from app.models.models import User, ProductionLine, CarModel, MasterData, ProductionData, MaterialData

EXCEL_PATH = os.path.join(os.path.dirname(__file__), 'RM_Master_Data_Final.xlsx')

# ─── 1. Users ────────────────────────────────────────────────────────────────
USERS_TO_SEED = [
    {"username": "admin",       "name": "Admin User",         "password": "admin",       "role": "Admin"},
    {"username": "manager",     "name": "Manoj Singh",        "password": "manager",     "role": "Manager"},
    {"username": "supervisor",  "name": "Rajesh Kumar",       "password": "supervisor",  "role": "Supervisor"},
    {"username": "deo",         "name": "Anil Verma",         "password": "deo",         "role": "DEO"},
    {"username": "ppc",         "name": "Priya Sharma",       "password": "ppc",         "role": "PPC_Planner"},
    {"username": "storekeeper", "name": "Suresh Patel",       "password": "storekeeper", "role": "Store_Keeper"},
]

# ─── 2. Production Lines ─────────────────────────────────────────────────────
PRODUCTION_LINES = [
    {"name": "LINE-A",  "description": "Main Assembly Line A"},
    {"name": "LINE-B",  "description": "Main Assembly Line B"},
    {"name": "LINE-C",  "description": "Sub Assembly Line C"},
    {"name": "320T",    "description": "320T Press Machine"},
    {"name": "400T",    "description": "400T Press Machine"},
    {"name": "500T",    "description": "500T Press Machine"},
    {"name": "630T",    "description": "630T Press Machine"},
    {"name": "800T",    "description": "800T Press Machine"},
]

# ─── 3. Car Models from Excel ─────────────────────────────────────────────────
MODEL_TYPES = {
    "BOLERO": "SUV",  "KUV": "SUV",     "XUV": "SUV",    "TUV": "MPV",
    "ARJUN": "Tractor", "LOADKING": "LCV", "ALFA": "3-Wheeler",
    "CURVV": "EV",    "EV": "EV",       "MPV": "MPV",    "COMPACT": "SUV",
    "TML": "LCV",     "MTBL": "Tractor","H3": "Tractor", "T4": "Tractor",
    "U202": "LCV",    "W601": "SUV",    "P112": "LCV",   "P113": "LCV",
    "Z101": "LCV",
}

def seed_users(session):
    print("\n[1/4] Seeding users...")
    created = 0
    for ud in USERS_TO_SEED:
        if not User.query.filter_by(username=ud["username"]).first():
            u = User(username=ud["username"], name=ud["name"], role=ud["role"])
            u.password = ud["password"]
            session.add(u)
            created += 1
            print(f"  [+] Created user: {ud['username']} ({ud['role']})")
        else:
            print(f"  [~] Exists: {ud['username']}")
    session.commit()
    print(f"  Users: {created} created.")

def seed_production_lines(session):
    print("\n[2/4] Seeding production lines...")
    created = 0
    for pl in PRODUCTION_LINES:
        if not ProductionLine.query.filter_by(name=pl["name"]).first():
            line = ProductionLine(name=pl["name"], description=pl["description"], is_active=True, status="AVAILABLE")
            session.add(line)
            created += 1
            print(f"  [+] Created line: {pl['name']}")
        else:
            print(f"  [~] Exists: {pl['name']}")
    session.commit()
    print(f"  Lines: {created} created.")

def seed_car_models(session, unique_models):
    print("\n[3/4] Seeding car models...")
    default_line = ProductionLine.query.first()
    supervisor = User.query.filter_by(role="Supervisor").first()
    deo = User.query.filter_by(role="DEO").first()
    created = 0

    for i, model_name in enumerate(sorted(unique_models)):
        code = f"{model_name[:3].upper()}-{i+1:02d}"
        if not CarModel.query.filter_by(name=model_name).first():
            cm = CarModel(
                name=model_name,
                model_code=code,
                type=MODEL_TYPES.get(model_name, "Standard"),
                production_line_id=default_line.id if default_line else None,
                supervisor_id=supervisor.id if supervisor else None,
                assigned_deo_id=deo.id if deo else None,
                status="ACTIVE",
            )
            session.add(cm)
            created += 1
            print(f"  [+] Created model: {model_name} ({code})")
        else:
            print(f"  [~] Exists: {model_name}")
    session.commit()
    print(f"  Models: {created} created.")

def seed_master_data(session):
    print("\n[4/4] Seeding master data from Excel...")
    if not os.path.exists(EXCEL_PATH):
        print(f"  [!] Excel file not found at {EXCEL_PATH}, skipping.")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb['RM Master Data']

    created = 0
    skipped = 0
    errors = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Col: SR_NO, MODEL, SAP_PART_NO, PART_NO, DESCRIPTION, ASSEMBLY_NO,
        #      RM_THK, SHEET_WIDTH, SHEET_LENGTH, COMP_PER_SHEET, RM_SIZE,
        #      RM_GRADE, ACT_RM_SIZES, REVISED, VALIDITY
        sr_no, model, sap_part, part_no, desc, assembly, \
            rm_thk, sheet_w, sheet_l, comp_per_sheet, rm_size, \
            rm_grade, act_rm, revised, validity = (list(row) + [None]*15)[:15]

        if not model or not sap_part:
            skipped += 1
            continue

        sap_part = str(sap_part).strip()
        model = str(model).strip()

        # Check for existing
        existing = MasterData.query.filter_by(model=model, sap_part_number=sap_part).first()
        if existing:
            skipped += 1
            continue

        try:
            md = MasterData(
                model=model,
                sap_part_number=sap_part,
                part_number=str(part_no).strip() if part_no else None,
                description=str(desc).strip() if desc else None,
                assembly_number=str(assembly).strip() if assembly else None,
            )
            session.add(md)
            session.flush()

            # Add material data with RM info
            if any([rm_thk, sheet_w, sheet_l, rm_size, rm_grade]):
                mat = MaterialData(
                    master_data_id=md.id,
                    additional_fields={
                        "RM Thk mm": str(rm_thk) if rm_thk is not None else None,
                        "Sheet Width": str(sheet_w) if sheet_w is not None else None,
                        "Sheet Length": str(sheet_l) if sheet_l is not None else None,
                        "No of comp per sheet": str(comp_per_sheet) if comp_per_sheet is not None else None,
                        "RM SIZE": str(rm_size) if rm_size is not None else None,
                        "RM Grade": str(rm_grade) if rm_grade is not None else None,
                        "Act RM Sizes": str(act_rm) if act_rm is not None else None,
                        "Revised": str(revised) if revised is not None else None,
                        "VALIDITY": str(validity) if validity is not None else None,
                    }
                )
                session.add(mat)

            created += 1
            if created % 50 == 0:
                session.commit()
                print(f"  ... {created} parts imported so far")

        except Exception as e:
            errors += 1
            session.rollback()
            print(f"  [!] Error on row {sr_no} ({model} / {sap_part}): {e}")

    session.commit()
    print(f"  Master data: {created} created, {skipped} skipped, {errors} errors.")

def run():
    print("=" * 60)
    print(" LPS System — Comprehensive Database Seeder")
    print("=" * 60)

    app = create_app('dev')
    with app.app_context():
        # Read unique models from Excel before seeding
        unique_models = set()
        if os.path.exists(EXCEL_PATH):
            wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
            ws = wb['RM Master Data']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1]:
                    unique_models.add(str(row[1]).strip())
            print(f"\nFound {len(unique_models)} unique models in Excel: {sorted(unique_models)}")

        seed_users(db.session)
        seed_production_lines(db.session)
        if unique_models:
            seed_car_models(db.session, unique_models)
        seed_master_data(db.session)

    print("\n" + "=" * 60)
    print(" ✓  Seeding complete! Restart backend and reload frontend.")
    print("=" * 60)

if __name__ == "__main__":
    run()
